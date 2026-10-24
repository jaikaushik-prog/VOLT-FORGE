"""Inventory FreeTwinEV metadata and CSV headers, without decoding measurement rows.

Source files are hashed as opaque bytes. Only README/LICENSE, the protocol
workbook and DOCX text are interpreted. No AllData XLSX workbook is opened.
"""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DOI = '10.5281/zenodo.18405243'


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024*1024), b''):
            digest.update(block)
    return digest.hexdigest()


def csv_header(path):
    # Binary first-line access prevents a text decoder buffering measurement rows.
    with Path(path).open('rb') as stream:
        line = stream.readline(65537)
    if len(line) > 65536 or not line.endswith(b'\n'):
        raise ValueError('Missing or oversized single-line CSV header: '+str(path))
    header = next(csv.reader([line.decode('utf-8-sig').strip()]))
    if not header or any(not x for x in header) or len(header) != len(set(header)):
        raise ValueError('Empty or duplicate CSV header fields: '+str(path))
    return header


def protocol_metadata(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = {sheet.title: dict(rows=sheet.max_row,columns=sheet.max_column,state=sheet.sheet_state) for sheet in book}
        # Metadata only, never a measurement file. Cell addresses preserve source attribution.
        cells = {sheet.title: {cell.coordinate:cell.value for row in sheet.iter_rows() for cell in row
                             if cell.value is not None} for sheet in book}
        matrix = book['Test_Matrix']
        tests = []
        for row in matrix.iter_rows(min_row=2,max_col=8):
            values = [c.value for c in row]
            if type(values[0]) is not int:
                continue
            tests.append(dict(id=values[0],timestamp=values[1],declared_discharge_A=values[2],
                declared_charge_A=values[3],declared_coolant_C=values[4],declared_flow_L_min=values[5],
                step7_DAQ=values[6],step8_DAQ=values[7],source_range=f'Test_Matrix!A{row[0].row}:H{row[0].row}'))
        mentions_calibration = any('Calibration_Table' in str(v) for values in cells.values() for v in values.values())
        return dict(sheets=sheets,cells=cells,tests=tests,
            calibration_sheet_referenced=mentions_calibration,calibration_sheet_present='Calibration_Table' in sheets)
    finally:
        book.close()


def inventory(source):
    source = Path(source).resolve()
    files = sorted(p for p in source.rglob('*') if p.is_file())
    before = {p.relative_to(source).as_posix():dict(bytes=p.stat().st_size,sha256=sha256(p)) for p in files}
    csv_files = sorted((source/'AllData').rglob('*.csv'))
    if not csv_files:
        raise ValueError('No measurement CSV headers found')
    records = []
    for path in csv_files:
        match = re.fullmatch(r'ID(\d+)_([0-9_]+)',path.parent.name)
        if not match:
            raise ValueError('Unknown measurement folder: '+path.parent.name)
        records.append(dict(id=int(match[1]),timestamp=match[2],folder=path.parent.name,
            csv=path.relative_to(source).as_posix(),xlsx=path.with_suffix('.xlsx').relative_to(source).as_posix(),
            xlsx_exists=path.with_suffix('.xlsx').is_file(),columns=csv_header(path),
            outcome_access='header only; no measurement rows interpreted',allocation='unallocated; complete record retained'))
    protocol = protocol_metadata(source/'Measurement_Protocol.xlsx')
    metadata_by_id = {x['id']:x for x in protocol['tests']}
    for record in records:
        record['protocol_matches'] = record['id'] in metadata_by_id and metadata_by_id[record['id']]['timestamp']==record['timestamp']
    document = next(source.glob('*.docx'))
    with ZipFile(document) as zipped:
        xml = ET.fromstring(zipped.read('word/document.xml'))
        ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        paragraphs = [''.join(x.text or '' for x in p.findall('.//w:t',ns)) for p in xml.findall('.//w:p',ns)]
        doc = dict(path=document.name,paragraphs=[p for p in paragraphs if p],
            embedded_media=[name for name in zipped.namelist() if '/media/' in name or '/embeddings/' in name])
    common = sorted(set.intersection(*(set(r['columns']) for r in records)))
    same = all(r['columns']==records[0]['columns'] for r in records)
    expected = ['timesamp','current','voltage','SOC','liquid_flow','temp_ambient','temp_inlet','temp_outlet']
    expected += [f'temp_TH{i}' for i in range(1,11)]
    expected += [f'temp_cell-{cell}_TH{sensor}' for cell,start in [(1,20),(2,17),(3,14),(4,11)] for sensor in range(start,start+3)]
    if not same or records[0]['columns'] != expected:
        raise ValueError('Schema differs from the inspected FreeTwinEV v1 release; review before making channel-coverage claims')
    after = {p.relative_to(source).as_posix():dict(bytes=p.stat().st_size,sha256=sha256(p)) for p in files}
    if before != after:
        raise ValueError('Source changed during inventory; discard this inventory and inspect the change')
    return dict(schema_version='voltforge.dataset_inventory/1',source_folder=source.name,
        citation=dict(doi=DOI,url='https://zenodo.org/records/18405243',version='v1',
            creators=['Milan Plzak','Simon Berta','Andrej Urge','Martin Bata'],license='CC BY 4.0',
            identity_basis='Local README/title and public record; archive-byte verification recorded separately'),
        file_count=len(files),total_bytes=sum(x['bytes'] for x in before.values()),files=before,
        record_count=len(records),records=records,uniform_headers=same,common_columns=common,
        protocol=protocol,document=doc,
        csv_xlsx_value_equivalence='not tested; requires opening measurement outcomes',
        outcome_rows_interpreted=0,measurement_xlsx_files_opened=0,model_fitting_performed=False,
        source_unchanged_during_inventory=True,
        validation_readiness=dict(frozen_LG_v2='unsupported: different cell model and unverified connection topology',
            branch_current_validation='unavailable: no individual branch current channels in the 30-column schema',
            cell_voltage_validation='unavailable: no individual cell voltage channels',
            current_conservation_check='unavailable: no measured branch-current sum',
            module_voltage_benchmark='candidate only; requires P45B-specific model, input/calibration audit and prospective allocation'),
        assessment_allocation='No development or assessment records selected; all complete records remain unopened')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,default=ROOT/'1s4p Battery Module Measurement')
    parser.add_argument('--out',type=Path,default=ROOT/'reports/freetwinev_inventory_v1/inventory.json')
    args = parser.parse_args()
    if args.out.resolve().is_relative_to(args.source.resolve()):
        raise SystemExit('Inventory must be outside original source directory')
    if args.out.exists():
        raise SystemExit('Choose a new output path to preserve existing evidence')
    result = inventory(args.source)
    args.out.parent.mkdir(parents=True,exist_ok=True)
    with args.out.open('x',encoding='utf-8',newline='\n') as stream:
        stream.write(json.dumps(result,indent=2,ensure_ascii=False,allow_nan=False)+'\n')
    print(f"Inventoried {result['file_count']} files and {result['record_count']} complete records; no measurement rows interpreted.")


if __name__=='__main__':
    main()
