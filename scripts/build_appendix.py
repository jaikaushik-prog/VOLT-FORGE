"""Build the Round 1 appendix workbook from the project's own records."""
from pathlib import Path
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'reports/appendix_v1/VOLT-FORGE_Appendix.xlsx'


def read(path): return json.loads((ROOT/path).read_text())


cost, market, parts = read('configs/build_cost_v1.json'), read('configs/market_sizing_v1.json'), read('configs/parts_prices_v1.json')
study, measured, protocol = read('reports/design_study_v2/run1/summary.json'), read('reports/measured_comparison_v1/final/result.json'), read('configs/measured_comparison_v1.json')
slide, spec = read('reports/slide_v4/data.json'), read('configs/round1_design_study_v2.json')

BLUE, DARK, GREY = '0B6FA8', '10202F', '5B6B7A'
INPUT, LINK = Font(name='Arial', size=10, color='0000FF'), Font(name='Arial', size=10, color='008000')
BODY, BOLD = Font(name='Arial', size=10, color=DARK), Font(name='Arial', size=10, bold=True, color=DARK)
ASSUMED, HEAD_FILL, BEST = PatternFill('solid', fgColor='FFFF00'), PatternFill('solid', fgColor=BLUE), PatternFill('solid', fgColor='DFF1FB')
WRAP = Alignment(wrap_text=True, vertical='top')
wb = Workbook()


def tab(name, title, line, widths):
    ws = wb.active if len(wb.worksheets) == 1 and wb.active.title == 'Sheet' else wb.create_sheet()
    ws.title, ws.sheet_view.showGridLines = name, False
    for i, w in enumerate(widths): ws.column_dimensions[chr(65+i)].width = w
    ws['A1'], ws['A1'].font = title, Font(name='Arial', size=15, bold=True, color=BLUE)
    ws['A2'], ws['A2'].font = line, Font(name='Arial', size=10, color=GREY)
    return ws


def section(ws, row, text):
    ws.cell(row, 1, text).font = Font(name='Arial', size=11, bold=True, color=BLUE)
    return row+1


def header(ws, row, names):
    for i, name in enumerate(names, 1):
        c = ws.cell(row, i, name); c.font, c.fill, c.alignment = Font(name='Arial', size=10, bold=True, color='FFFFFF'), HEAD_FILL, Alignment(wrap_text=True, vertical='center')
    return row+1


def put(ws, row, col, value, kind='body', fmt=None, assumed=False):
    c = ws.cell(row, col, value)
    c.font = {'body': BODY, 'bold': BOLD, 'input': INPUT, 'link': LINK}[kind]; c.alignment = Alignment(vertical='top', wrap_text=isinstance(value, str) and not value.startswith('='))
    if fmt: c.number_format = fmt
    if assumed: c.fill = ASSUMED
    return c


def aside(ws, row, col, text):
    put(ws, row, col, text).alignment = Alignment(vertical='top', indent=1)


# Build cost
b = tab('Build cost', 'What the full product would cost to build', 'About ₹35 crore over three years. Estimate, not a quote. Change any blue cell and everything follows.', [46, 16, 12, 12, 12, 12, 12, 12, 12, 12])
r = section(b, 4, 'Inputs')
inputs = {}
for label, key, fmt, why in [('Overhead on salary', 'overhead_multiplier', '0.0"×"', 'Employer costs, hardware, office, hiring, attrition'), ('Contingency', 'contingency', '0%', 'A research-heavy build carries schedule risk'), ('₹ per US$', 'usd_inr', '0', 'Planning rate')]:
    put(b, r, 1, label); put(b, r, 2, cost[key]['value'], 'input', fmt, True); aside(b, r, 3, why); inputs[key] = f'$B${r}'; r += 1
put(b, r, 1, 'Duration, years'); put(b, r, 2, cost['duration_years'], 'input', '0', True); inputs['years'] = f'$B${r}'; r += 2
r = section(b, r, 'Work breakdown, in person-years')
roles = ['principal', 'senior_physics', 'senior_software', 'engineer', 'test_technician']
r = header(b, r, ['Workstream', 'Vision section', 'Principal', 'Senior physics', 'Senior software', 'Engineer', 'Test technician', 'Person-years', '₹ lakh', '₹ crore'])
put(b, r, 1, 'Pay, ₹ lakh a year', 'bold')
for i, role in enumerate(roles): put(b, r, 3+i, cost['rates_lakh_per_year'][role]['value'], 'input', '0.0')
rate_row, first = r, r+1; r += 1
for item in cost['work_breakdown_person_years']:
    put(b, r, 1, item['name']); put(b, r, 2, item['vision'])
    for i, role in enumerate(roles): put(b, r, 3+i, item.get(role, 0), 'input', '0.0;-0.0;-')
    put(b, r, 8, f'=SUM(C{r}:G{r})', fmt='0.0'); put(b, r, 9, f'=SUMPRODUCT($C${rate_row}:$G${rate_row},C{r}:G{r})*{inputs["overhead_multiplier"]}', fmt='#,##0.0'); put(b, r, 10, f'=I{r}/100', fmt='0.00'); r += 1
last = r-1
put(b, r, 1, 'Total people', 'bold')
for col in 'HIJ': put(b, r, ord(col)-64, f'=SUM({col}{first}:{col}{last})', 'bold', {'H': '0.0', 'I': '#,##0.0', 'J': '0.00'}[col])
people_row = r; r += 2
r = section(b, r, 'Lab')
r = header(b, r, ['Item', 'Basis', 'Quantity', 'Unit ₹', 'Unit US$', '₹ lakh'])
lab_first = r
for item in cost['lab']:
    is_assumed = item['basis'] == 'assumed'
    put(b, r, 1, item['item']); put(b, r, 2, item['basis']); put(b, r, 3, item['quantity'], 'input', '#,##0'); put(b, r, 4, item.get('unit_inr'), 'input', '#,##0', is_assumed); put(b, r, 5, item.get('unit_usd'), 'input', '#,##0')
    put(b, r, 6, f'=C{r}*IF(D{r}<>"",D{r},E{r}*{inputs["usd_inr"]})/100000', fmt='#,##0.0'); r += 1
put(b, r, 1, 'Total lab', 'bold'); put(b, r, 6, f'=SUM(F{lab_first}:F{r-1})', 'bold', '#,##0.0'); lab_row = r; r += 2
r = section(b, r, 'Running costs (all assumed)')
r = header(b, r, ['Item', 'Basis', '₹ lakh a year', '₹ lakh total'])
run_first = r
for item in cost['recurring_lakh_per_year']:
    put(b, r, 1, item['item']); put(b, r, 2, 'assumed'); put(b, r, 3, item['value'], 'input', '0', True); put(b, r, 4, f'=C{r}*{inputs["years"]}', fmt='#,##0'); r += 1
put(b, r, 1, 'Total running', 'bold'); put(b, r, 4, f'=SUM(D{run_first}:D{r-1})', 'bold', '#,##0'); run_row = r; r += 1
put(b, r, 1, cost['one_off_lakh'][0]['item']+' (one-off)'); put(b, r, 2, 'assumed'); put(b, r, 4, cost['one_off_lakh'][0]['value'], 'input', '#,##0', True); one_row = r; r += 2
r = section(b, r, 'Total')
r = header(b, r, ['Layer', '₹ crore', 'Share'])
total_first = r
for label, formula in [('People', f'=I{people_row}/100'), ('Lab', f'=F{lab_row}/100'), ('Running costs', f'=D{run_row}/100'), ('One-off', f'=D{one_row}/100')]:
    put(b, r, 1, label); put(b, r, 2, formula, fmt='0.00'); r += 1
put(b, r, 1, 'Contingency'); put(b, r, 2, f'=SUM(B{total_first}:B{r-1})*{inputs["contingency"]}', fmt='0.00'); r += 1
put(b, r, 1, 'Total build', 'bold'); put(b, r, 2, f'=SUM(B{total_first}:B{r-1})', 'bold', '0.00'); total_row = r
for row in range(total_first, r+1): put(b, row, 3, f'=B{row}/$B${total_row}', fmt='0%')
r += 1
put(b, r, 1, 'In US$ million'); put(b, r, 2, f'=B{total_row}*10/{inputs["usd_inr"]}', fmt='0.0'); usd_row = r; r += 1
put(b, r, 1, 'Average team size'); put(b, r, 2, f'=H{people_row}/{inputs["years"]}', fmt='0.0'); team_row = r; r += 2
r = section(b, r, 'Cross-check: COCOMO II, a standard software estimate')
ksloc_first = r
for key, value in cost['cross_check']['ksloc'].items():
    put(b, r, 1, 'Thousand lines: '+key.replace('_', ' ')); put(b, r, 2, value, 'input', '0', True); r += 1
put(b, r, 1, 'Thousand lines, total', 'bold'); put(b, r, 2, f'=SUM(B{ksloc_first}:B{r-1})', 'bold', '0'); ks = r; r += 1
put(b, r, 1, 'Person-months = 2.94 × lines^1.0997'); put(b, r, 2, f'=2.94*B{ks}^1.0997', fmt='#,##0'); pm = r; r += 1
put(b, r, 1, 'Cost at our average pay, ₹ crore'); put(b, r, 2, f'=B{pm}/12*I{people_row}/H{people_row}/100', fmt='0.00'); cc = r; r += 1
put(b, r, 1, 'Difference from our people cost'); put(b, r, 2, f'=B{cc}/J{people_row}-1', fmt='0%'); aside(b, r, 3, 'Two methods land close. Software lines only, so a lower number is expected.'); r += 2
r = section(b, r, 'Benchmark: AspenTech')
put(b, r, 1, 'Research and development, FY2024, US$ million'); put(b, r, 2, cost['benchmark']['usd_million'], 'input', '0.0'); bm = r; r += 1
put(b, r, 1, 'In ₹ crore'); put(b, r, 2, f'=B{bm}*{inputs["usd_inr"]}/10', fmt='#,##0'); bmc = r; r += 1
put(b, r, 1, 'Our whole build as a share of one AspenTech year', 'bold'); put(b, r, 2, f'=B{total_row}/B{bmc}', 'bold', '0.0%'); share_row = r

# Market
m = tab('Market', 'Who could use it, India first', 'Students use it free, so revenue comes from labs and teams. Counts size the audience; they do not show anyone will pay.', [50, 16, 62])
s, a = market['sourced'], market['assumed']
r = section(m, 4, 'From AISHE 2021-22 (Ministry of Education)')
r = header(m, r, ['Group', 'Count', 'Note'])
rows = {}
for key, label, note in [('engineering_technology_enrolment', 'Engineering and technology students', ''), ('electronics_engineering_enrolment', 'Electronics students', ''), ('mechanical_engineering_enrolment', 'Mechanical students', ''), ('electrical_engineering_enrolment', 'Electrical students', ''), ('engineering_technology_colleges', 'Engineering and technology colleges', market['source']['colleges_base'])]:
    put(m, r, 1, label); put(m, r, 2, s[key], 'input', '#,##0'); put(m, r, 3, note); rows[key] = r; r += 1
r += 1; r = section(m, r, 'Assumed, nobody has seen this price')
put(m, r, 1, 'Lab licence, ₹ a year'); put(m, r, 2, a['annual_lab_licence_inr']['value'], 'input', '#,##0', True); put(m, r, 3, a['annual_lab_licence_inr']['reason']); price = r; r += 1
put(m, r, 1, 'Share of colleges reached'); put(m, r, 2, a['first_stage_college_share']['value'], 'input', '0%', True); put(m, r, 3, a['first_stage_college_share']['reason']); reach = r; r += 2
r = section(m, r, 'Result')
r = header(m, r, ['Level', 'Value', 'What it is'])
put(m, r, 1, 'TAM, students (lakh)', 'bold'); put(m, r, 2, f'=B{rows["engineering_technology_enrolment"]}/100000', 'bold', '0.0'); put(m, r, 3, market['funnel']['tam']); tam = r; r += 1
put(m, r, 1, 'SAM, students (lakh)', 'bold'); put(m, r, 2, f'=SUM(B{rows["electronics_engineering_enrolment"]}:B{rows["electrical_engineering_enrolment"]})/100000', 'bold', '0.0'); put(m, r, 3, market['funnel']['sam']); sam = r; r += 1
put(m, r, 1, 'SOM, colleges', 'bold'); put(m, r, 2, f'=ROUNDDOWN(B{rows["engineering_technology_colleges"]}*B{reach},0)', 'bold', '0'); put(m, r, 3, market['funnel']['som']+', rounded down'); som = r; r += 1
put(m, r, 1, 'SOM revenue, ₹ lakh a year'); put(m, r, 2, f'=B{som}*B{price}/100000', fmt='0.0'); rev = r; r += 2
r = section(m, r, 'Honest finding')
put(m, r, 1, 'Years of SOM revenue to repay the build', 'bold'); put(m, r, 2, f"='Build cost'!B{total_row}*100/B{rev}", 'link', '0'); put(m, r, 3, 'Lab licences alone will not repay a ₹35 crore build. The plan needs a paid tier for engineering teams. That price is not set yet.'); payback = r

# Parts
p = tab('Parts', 'What the backup would cost in parts', 'Public listings checked 30 Oct 2026. One exact part per line. Two cells halve the cell bill, not the whole unit.', [34, 16, 12, 10, 16, 24, 54])
r = section(p, 4, 'Listings')
r = header(p, r, ['Part', 'Seller', 'Price', 'Currency', 'Tax', 'Availability', 'Note'])
lines = {}
for key in ['cell:LG_M50T_ECN', 'cell:SAMSUNG_30T', 'copper', 'converter']:
    item, l = parts['items'][key], parts['items'][key]['listings'][0]
    put(p, r, 1, item['name']); put(p, r, 2, l['seller']); put(p, r, 3, l['price'], 'input', '#,##0.00'); put(p, r, 4, l['currency']); put(p, r, 5, l['tax']); put(p, r, 6, l['availability']); put(p, r, 7, ' '.join(filter(None, [item.get('exclusions'), item.get('limits')]))); lines[key] = r; r += 1
put(p, r, 1, parts['items']['protection']['name']); put(p, r, 7, parts['items']['protection']['note']+'. '+parts['items']['protection']['specification']); r += 2
r = section(p, r, 'Cost by cell count')
r = header(p, r, ['Cells', 'Mass, g', 'Cells, ₹', 'Board, ₹', 'Cells + board, ₹', 'Saving against 4 cells', 'Note'])
put(p, r+4, 1, 'Mass per cell, g'); put(p, r+4, 2, 70, 'input', '0'); per = r+4
base = r
for n in range(1, 5):
    put(p, r, 1, n, 'bold' if n == 2 else 'body', '0'); put(p, r, 2, f'=A{r}*$B${per}', fmt='0'); put(p, r, 3, f'=A{r}*$C${lines["cell:LG_M50T_ECN"]}', fmt='#,##0')
    put(p, r, 4, f"=$C${lines['converter']}*'Build cost'!{inputs['usd_inr']}", 'link', '#,##0'); put(p, r, 5, f'=C{r}+D{r}', fmt='#,##0'); put(p, r, 6, f'=1-E{r}/$E${base+3}', fmt='0%')
    put(p, r, 7, {1: 'Refused: over the cell current limit', 2: 'Recommended', 3: '', 4: 'Starting design'}[n])
    if n == 2:
        for col in range(1, 8): p.cell(r, col).fill = BEST
    r += 1
put(p, per+1, 1, 'Cells + board leaves out shipping, import duty, links and enclosure. The board is not confirmed to run with fewer than four cells.'); p.merge_cells(start_row=per+1, start_column=1, end_row=per+1, end_column=7)

# Design study
d = tab('Design study', 'How few cells can back up a Raspberry Pi 5?', 'Two. 4,006 simulated cases, frozen before the first run. Every row is a design; the blue row is the recommendation.', [10, 26, 11, 10, 10, 11, 11, 11, 11, 11, 11, 11])
r = section(d, 4, 'The job')
job = {}
for key, label, value, fmt, assumed in [('w', 'Pi load, W', spec['missions']['pi']['power_W'], '0', False), ('e', 'Converter efficiency', spec['missions']['pi']['converter_efficiency'], '0%', True), ('v', 'Voltage floor, V', spec['common']['voltage_floor_V'], '0.0', True), ('lim', 'Cell current limit, A', 7.275, '0.000', False)]:
    put(d, r, 1, label); d.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2); put(d, r, 3, value, 'input', fmt, assumed); job[key] = f'$C${r}'; r += 1
put(d, r, 1, 'Battery current held, A'); d.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2); put(d, r, 3, f'={job["w"]}/{job["e"]}/{job["v"]}', fmt='0.00'); aside(d, r, 4, 'Constant-current stand-in for a constant-power load, for 10 minutes at 23 °C'); r += 2
r = section(d, r, 'Results')
r = header(d, r, ['Cells', 'Output layout', 'Verdict', 'Required passed', 'Required total', 'Lowest V, normal', 'Peak A, normal', 'Lowest V, worst', 'Peak A, worst', 'Worst °C', 'Spare A', '12 A stress samples passed'])
names = {'baseline': 'Same end', 'opposite': 'Far cell', 'thicker': 'Same end, thicker links', 'combined': 'Far cell, thicker links'}
top = r
for x in study['designs']:
    w, n = x['worst_required'], x['nominal']; passed = x['qualifies']
    put(d, r, 1, x['count'], fmt='0'); put(d, r, 2, names[x['option']]); put(d, r, 3, 'Qualifies' if passed else 'Refused'); put(d, r, 4, x['required_passed'], fmt='0'); put(d, r, 5, x['required_cases'], fmt='0')
    if passed:
        for col, value, fmt in [(6, n['min_voltage_V'], '0.00'), (7, n['peak_cell_current_A'], '0.00'), (8, w['lowest_voltage_V'], '0.00'), (9, w['highest_cell_current_A'], '0.00'), (10, w['highest_temperature_C'], '0.0'), ]: put(d, r, col, value, fmt=fmt)
        put(d, r, 11, f'={job["lim"]}-I{r}', fmt='0.00'); put(d, r, 12, f"{x['samples']['stress']['passed']} of {x['samples']['stress']['cases']}").alignment = Alignment(horizontal='right')
    if x['count'] == study['recommended']['count'] and x['option'] == study['recommended']['option']:
        for col in range(1, 13): d.cell(r, col).fill = BEST
        best = r
    r += 1
end = r-1; r += 1
r = section(d, r, 'The study at a glance')
facts = [('Cases run', study['cases'], '#,##0'), ('Refused by the product, all for going over the cell current limit', study['refused'], '0'), ('Calculated cases, every one closed its energy balance', f'={study["cases"]}-{study["refused"]}', '#,##0'), ('Time-step check: largest voltage change when the step was halved, V', study['half_step']['max_voltage_change_V'], '0.00000'), ('Time-step check: largest temperature change, K', study['half_step']['max_temperature_change_K'], '0.0000')]
for label, value, fmt in facts:
    put(d, r, 1, label); d.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8); put(d, r, 9, value, fmt=fmt); r += 1
r += 1; r = section(d, r, 'What argues against two cells')
for text in ['Less margin: 5.33 A worst case, against 3.34 A for four cells.', 'The 12 A stress test can fail: 2 of 62 mixed samples were refused. All 62 passed at the Pi load.', 'No link layout wins everywhere: the far-cell output raised peak current in 35 of 81 paired cases.', 'Temperatures are a range, not a prediction. See Evidence.']:
    put(d, r, 1, text); d.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12); r += 1

# Evidence
e = tab('Evidence', 'How far the results can be trusted', 'The maths checks out. Real-data comparison did not: all 12 cases missed. We kept that result and did not tune the model.', [34, 30, 12, 12, 12, 12, 12, 12, 11, 11])
r = section(e, 4, 'Checks')
r = header(e, r, ['Check', 'Result', 'What it means']); e.merge_cells(start_row=r-1, start_column=3, end_row=r-1, end_column=10)
check_first = r
checks = [
    ('Python tests', '438 passed, 8 known skips', 'Clean install; one test added later also passes'),
    ('JavaScript tests', '50 passed', 'Editor and app logic'),
    ('Independent solver (PyBaMM)', f"{slide['pybamm']['passed']} of {slide['pybamm']['cases']} cases", 'Same answer on voltage and heat, so the arithmetic is sound. This is not a test against real cells'),
    ('Energy balance', 'Closed in all 3,964 runs', 'Energy out of the cells equals energy delivered plus heat'),
    ('Export replay', f"{slide['replay']['passed']} of {slide['replay']['cases']} replayed", 'A saved test reproduces exactly; tampered files are refused'),
    ('Data files', f"{slide['kpi']['data_files']} files hashed", f"{slide['kpi']['lab_experiments']} module experiments from Piombo et al., inventory and integrity only"),
    ('Sealed records', f"{slide['kpi']['sealed_total']} never opened", f"{slide['kpi']['sealed_module_records']} module tests and {slide['kpi']['sealed_cell_records']} Samsung cell records, reserved for a final assessment"),
    ('Audit of V3 findings', '15 fixed, 3 partly, 2 open', 'Self-check, not independent. Open: VF3-08, VF3-19'),
]
for name, result, meaning in checks:
    put(e, r, 1, name); put(e, r, 2, result); put(e, r, 3, meaning); e.merge_cells(start_row=r, start_column=3, end_row=r, end_column=10); r += 1
r += 1; r = section(e, r, 'Compared with published measurements: frozen before any run, never used to fit')
v, t = protocol['bands']['voltage'], protocol['bands']['temperature']
put(e, r, 1, f"Voltage band: RMSE up to {v['rmse_V']*1000:.0f} mV and largest error up to {v['max_abs_V']*1000:.0f} mV. Temperature band: peak rise error up to 1 K or 25% of the measured rise, whichever is larger."); e.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10); e.row_dimensions[r].height = 28; r += 1
r = header(e, r, ['Case', 'Role', 'Current, A', 'Voltage error, V', 'Temp error, K', 'Rise measured, K', 'Rise simulated, K', 'Difference, K', 'Voltage in band', 'Temp in band'])
labels = {'primary': 'Exact M50T module', 'supporting': 'Module, already used for cooling', 'independent_adjacent_cell': 'Warwick M50 single cell'}
mt = r
for c in measured['cases']:
    name = c['case'].replace('test28_exact_m50t_module', 'Test 28').replace('test53_used_fixture', 'Test 53').replace('warwick_lg_m50_single_cell__LGM50_', 'Warwick ').replace('__hall_sum_current', ', Hall current').replace('__link_plus_1_mohm', ', link +1 mΩ').replace('__', ' ')
    put(e, r, 1, name); put(e, r, 2, labels[c['role']]); put(e, r, 3, c['current_A'], fmt='0.0'); put(e, r, 4, c['voltage_rmse_V'], fmt='0.000'); put(e, r, 5, c['temperature_rmse_K'], fmt='0.00'); put(e, r, 6, c['peak_rise_measured_K'], fmt='0.00'); put(e, r, 7, c['peak_rise_simulated_K'], fmt='0.00')
    put(e, r, 8, f'=G{r}-F{r}', fmt='+0.00;-0.00'); put(e, r, 9, 'Yes' if c['voltage_within_band'] else 'No').alignment = Alignment(horizontal='center'); put(e, r, 10, 'Yes' if c['temperature_within_band'] else 'No').alignment = Alignment(horizontal='center'); r += 1
put(e, r, 1, 'Cases inside both bands', 'bold'); put(e, r, 9, f'=COUNTIFS(I{mt}:I{r-1},"Yes",J{mt}:J{r-1},"Yes")', 'bold', '0'); put(e, r, 10, f'=ROWS(I{mt}:I{r-1})', 'bold', '0'); inband = r; r += 1
put(e, r, 1, 'What this means: single cells heat up more than the model says, and the test rig ran cooler than it says. So cooling depends on the rig, and every temperature stays a range. Voltage also stayed outside its band.'); e.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10); e.row_dimensions[r].height = 42

# Limits and sources
l = tab('Limits and sources', 'What is assumed, untested and where each number came from', 'Anything marked assumed is our judgement. No user, price or accuracy claim here comes from a real trial.', [38, 22, 86])
r = section(l, 4, 'Not done yet')
for text in ['Nothing has been built. All results are simulated.', 'No hardware test and no real user trial. Both are the next step.', 'Not tested: Windows, browsers other than Chrome, the Codex AI route.', 'No independent audit yet.', 'The UPS board may not run with fewer than four cells, and publishes no cutoff or efficiency.', 'The LG low-charge table change is recorded as unvalidated.']:
    put(l, r, 1, text); l.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); r += 1
r += 1; r = section(l, r, 'Sources')
r = header(l, r, ['Number', 'Basis', 'Source'])
src = [(k.replace('_', ' ').capitalize()+f" pay, ₹{v['value']} lakh", v['basis'], v['source']+(' '+v['url'] if v.get('url') else '')) for k, v in cost['rates_lakh_per_year'].items()]
src += [(i['item'], i['basis'], (i.get('source') or i.get('reason') or 'Our planning figure')+(' '+i['url'] if i.get('url') else '')) for i in cost['lab']]
src += [('COCOMO II formula', 'sourced', cost['cross_check']['source']), ('AspenTech R&D', 'sourced', cost['benchmark']['source']+' '+cost['benchmark']['url']), ('Students and colleges', 'sourced', market['source']['name']+' '+market['source']['url']+' (read from '+market['source']['read_from']+')')]
src += [(parts['items'][k]['name'], 'sourced', parts['items'][k]['listings'][0]['url']) for k in ['cell:LG_M50T_ECN', 'cell:SAMSUNG_30T', 'copper', 'converter']]
src += [('Module test data', 'sourced', 'Piombo et al., DOI 10.17632/zh58byr53c.2, CC BY 4.0'), ('Single-cell test data', 'sourced', 'Chen et al., Warwick LG M50 data, https://wrap.warwick.ac.uk/id/eprint/149527')]
for row in src:
    for col, value in enumerate(row, 1): put(l, r, col, value)
    r += 1

# Summary
u = tab('Summary', 'VOLT-FORGE appendix', 'The numbers behind the Round 1 slide. Tata Elxsi Teliport Season 4, Case Study 2. Results are simulated; nothing has been built.', [46, 16, 62])
r = section(u, 4, 'Headlines')
r = header(u, r, ['Number', 'Value', 'What it means'])
heads = [
    ('Smallest design that qualifies, cells', f"='Design study'!A{best}", '0', 'Two LG M50T cells with the far-cell output, among the parallel layouts tested'),
    ('Required Pi tests passed', f"='Design study'!D{best}", '0', 'Out of 33: normal, weak cell in each position, low cooling and more'),
    ('Worst cell current, A', f"='Design study'!I{best}", '0.00', 'Against a 7.275 A limit'),
    ('Cell cost for 2 cells, ₹', f"=Parts!C{base+1}", '#,##0', 'Half the ₹4,824 of the four-cell start. Cells only'),
    ('Simulated cases', f"='Design study'!I{top+len(study['designs'])+2}", '#,##0', 'Frozen plan, run once, every result kept'),
    ('Measured cases inside both bands', f"=Evidence!I{inband}", '0', 'Out of 12. The model was not tuned to close this gap'),
    ('Cost to build the full product, ₹ crore', f"='Build cost'!B{total_row}", '0.0', 'Three years, from own lab data to a full engineering app'),
    ('Person-years', f"='Build cost'!H{people_row}", '0.0', 'About 34 people on average'),
    ('Share of one AspenTech R&D year', f"='Build cost'!B{share_row}", '0.0%', 'Scale check against an established engineering software company'),
    ('TAM, students (lakh)', f"=Market!B{tam}", '0.0', 'All engineering and technology students in India'),
    ('SAM, students (lakh)', f"=Market!B{sam}", '0.0', 'Electrical, electronics and mechanical'),
    ('SOM, colleges', f"=Market!B{som}", '0', 'Assumed 5% of engineering colleges'),
]
for label, formula, fmt, meaning in heads:
    put(u, r, 1, label); put(u, r, 2, formula, 'link', fmt); aside(u, r, 3, meaning); r += 1
r += 1; r = section(u, r, 'Tabs')
for name, text in [('Build cost', 'Every workstream, salary, lab item and running cost, plus two cross-checks'), ('Market', 'TAM, SAM, SOM and how long licences alone take to repay the build'), ('Parts', 'Listings and cost at each cell count'), ('Design study', 'All 13 designs and what argues against the answer'), ('Evidence', 'Software checks and the comparison with real measurements'), ('Limits and sources', 'What is untested, and where each number came from')]:
    put(u, r, 1, name, 'bold'); put(u, r, 2, text); u.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3); r += 1
r += 1; r = section(u, r, 'Colour key')
put(u, r, 1, 'Blue text', 'input'); put(u, r, 2, 'An input you can change'); u.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3); r += 1
put(u, r, 1, 'Black text'); put(u, r, 2, 'A formula or a recorded result'); u.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3); r += 1
put(u, r, 1, 'Green text', 'link'); put(u, r, 2, 'Pulled from another tab'); u.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3); r += 1
put(u, r, 1, 'Yellow fill', assumed=True); put(u, r, 2, 'Our assumption, not a sourced number'); u.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)

for row in m.iter_rows(min_row=5, min_col=3, max_col=3):
    if row[0].value and not row[0].fill.fgColor.rgb.endswith(BLUE): row[0].alignment = Alignment(vertical='top', wrap_text=True, indent=1)
wb.move_sheet('Summary', offset=-len(wb.worksheets)+1)
wb.active = 0
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
