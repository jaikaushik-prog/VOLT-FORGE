"""Place the rendered slide on the official Teliport template, with sources in the notes."""
from pathlib import Path
import argparse
import json

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]


TEAM = dict(name='StratWise', members=['Shubhra Arora', 'Jai Kaushik', 'Sanya Bansal'], size_pt=18)


def notes(d):
    best, stress, pi = d['best'], d['best']['samples']['stress'], d['best']['samples']['pi']
    sh, tb, vb = d['checks']['post_hoc']['share_detail'][0], d['checks']['error_budget']['temperature'], d['checks']['error_budget']['voltage']
    return f"""Simulation only. Nothing has been built or measured.

Result: {d['study']['claim']}. Recommended: {best['count']} LG M50T cells with the output at the far cell. {best['required_passed']} of {best['required_cases']} required Pi cases passed, worst-case cell current {best['worst_required']['highest_cell_current_A']:.2f} A against a 7.275 A limit.
Limits: {stress['cases']-stress['passed']} of {stress['cases']} mixed samples were refused at the 12 A stress test; all {pi['cases']} passed at the Pi load. No link option won every case.
Mission: 25 W from the Raspberry Pi 5 supply recommendation, an assumed 90% converter and an assumed 3.3 V floor give 8.42 A, held constant as an approximation. Run time and room temperature are assumptions.
Real data: all {d['proof_cases']} pre-declared full-discharge checks missed their bands, on voltage and on temperature. These are 8 measured discharges and 4 sensitivity reruns. Sources: Chen et al. 2020 (LG M50, zenodo.4032561) and Piombo et al. 2024 test 28 (M50T, 10.17632/zh58byr53c.2).
Board: Geekworm X1205 2-cell UPS HAT, US$46, 5.1 V at up to 6 A, up to 95% efficiency and a 2.4 V cutoff, from its listing and wiki. Its battery overcurrent threshold is not published.
Current shares: Piombo test 28, first 10% of the discharge, measured against simulated, with Hall channels normalised to their own sum. Average error {sh['model_error_points']:.2f} points against {sh['equal_share_error_points']:.2f} for an equal-share guess, inside the frozen 2.5 point band. The busiest cell was named in {100*d['checks']['current_sharing'][0]['busiest_cell_match_fraction']:.0f}% of samples against a required 80%, so that part missed. The equal-share comparison was added after the run. This is one test with 3 mOhm links, early in the discharge; later on the error is larger, and the near-zero-link test shared evenly. Over the first 10 minutes, real heating was {min(x['ratio'] for x in d['checks']['post_hoc']['ten_minute_rise']):.2f} to {max(x['ratio'] for x in d['checks']['post_hoc']['ten_minute_rise']):.2f} times the model across 8 discharges, also added after the run.
Heat: real rises reached {tb['largest_measured_over_simulated']:.2f} times the simulated rise. Applied to the 2-cell worst case that gives {tb['scaled_C']:.1f} C against the LG 55 C limit. Voltage: the largest error in the Pi's charge range was {vb['primary']['largest_abs_error_V']:.2f} V against a {vb['margin_V']:.2f} V margin.
Not proven: a real Pi draw on this board, the board's overcurrent limit, or measured accuracy. Below 10 C the LG limit falls to 2.425 A a cell, so two cells could not carry this load.
Cost: cells only, Robu LG M50T Rs 1,206 including GST, listed out of stock, checked 30 Oct 2026. Excludes shipping and fabrication.
Market: AISHE 2021-22. The price of Rs {d['market']['licence_inr']:,} a lab and 5% reach are assumptions, and counts do not show willingness to pay.
Business: free for students, labs and teams pay, as a hypothesis. Stage one is dataset, results, quality and the lab, Rs {d['business']['stage1_crore']:.1f} crore. That is Rs {d['business']['yearly_crore']:.1f} crore a year for {d['business']['repay_years']} years, before running costs. {d['business']['labs']} labs at Rs 50,000 give Rs 65 lakh, and {d['business']['stage1_teams']} teams at an assumed Rs 5 lakh give the rest. No customer has seen these prices, and the number of paying teams has not been sized. One seat of MATLAB, Simulink, Simscape and Simscape Battery lists at US$13,180 perpetual, September 2026.
Build cost: about Rs {d['cost']['total_crore']:.0f} crore over 3 years and {d['cost']['person_years']:g} person-years for the full product vision. An estimate from sourced salary rates and assumed staffing, cross-checked with COCOMO II. See reports/build_cost_v1/BUILD_COST.md. AspenTech fiscal 2024 R&D was US$206.1 million.
Safety and limits: no spare cell, so three cells is the smallest design that survives losing one. 25 W is the supply rating, and below about 21.6 W one cell stays under its limit. Bench safety steps are in the appendix. No user has tried the app yet.
KPIs: {d['kpi']['lab_experiments']} module experiments in the Piombo et al. dataset, and {d['kpi']['data_files']} dataset files checksum-verified, including cell characterisation and calibration files. PyBaMM, an independent solver, matched voltage and heat in {d['pybamm']['passed']} of {d['pybamm']['cases']} cases. That shows the maths is sound, not that the model matches real cells. {d['kpi']['sealed_total']} records stay closed for a later exam.
Link: the site shows the demo video first, then the real engine running in the visitor's browser. The AI uses Gemini through a small server function.
Screens: real designs in the app. The table compares the 4-cell starting design with the 2-cell design under the same 8.42 A load."""


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--template', type=Path, required=True)
    parser.add_argument('--out', type=Path, default=ROOT/'reports/slide_v4/VOLT-FORGE_Round1_Slide.pptx')
    args = parser.parse_args()
    d = json.loads((ROOT/'reports/slide_v4/data.json').read_text()); prs = Presentation(args.template)
    cover, slide = prs.slides[0], prs.slides[1]
    for shape in cover.shapes:
        if shape.has_text_frame and 'Case Study No' in shape.text_frame.text:
            shape.text_frame.paragraphs[0].runs[0].text = 'Case Study 2'
            for extra in shape.text_frame.paragraphs[0].runs[1:]: extra.text = ''
        if shape.has_text_frame and 'Team Details' in shape.text_frame.text:
            frame = shape.text_frame; frame.paragraphs[0].runs[0].text = f"Team {TEAM['name']}"
            for extra in frame.paragraphs[0].runs[1:]: extra.text = ''
            line = frame.add_paragraph(); line.text = ', '.join(TEAM['members']); line.font.size = Pt(TEAM['size_pt']); line.font.bold = False
    for shape in list(slide.placeholders):
        if shape.placeholder_format.idx != 0: shape._element.getparent().remove(shape._element)
    slide.shapes.title.text_frame.text = 'Test a battery before you build it'
    picture = slide.shapes.add_picture(str(ROOT/'reports/slide_v4/VOLT-FORGE_Round1_Slide.png'), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    picture.name = 'VOLT-FORGE Round 1 slide'
    links = json.loads((ROOT/'configs/public_site.json').read_text()); boxes = json.loads((ROOT/'reports/slide_v4/link_boxes.json').read_text())
    for key, address in (('site', 'https://'+links['address']), ('appendix', f"https://{links['address']}/appendix")):
        x, y, w, h = [Pt(v/2) for v in boxes[key]]  # The page is 1920 px wide on a 960 pt slide.
        spot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h); spot.name = key+' link'; spot.fill.background(); spot.line.fill.background(); spot.click_action.hyperlink.address = address
    picture._element.nvPicPr.cNvPr.set('descr', 'VOLT-FORGE: the seven rubric criteria around the working app, showing the 4-cell Pi circuit with a cell properties panel, and a comparison of the 4-cell and 2-cell designs.')
    slide.notes_slide.notes_text_frame.text = notes(d)
    prs.save(args.out)


if __name__ == '__main__':
    main()
